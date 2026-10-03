"""End-to-end API regression suite for Фактура+.

Committed version of the ad-hoc manual test script used throughout
development, converted to real pytest so it runs as part of CI/local
`pytest` instead of only ever being run by hand. Tests share one
in-memory database and largely one owner/staff pair across the whole
module (like the original script), so they run in file order and build
on each other's state - this mirrors real usage (an owner invites staff,
enters revenue, then checks permission gating against that same data)
rather than testing each endpoint in total isolation.
"""
import time
from datetime import datetime

from .conftest import auth_headers as h

TS = int(time.time())
CURRENT_MONTH = datetime.now().strftime("%Y-%m")

# Shared across test functions within this module, mirroring the shared
# variables the original manual script used (owner_token, invoice_id, ...).
STATE = {}


def test_auth(client):
    owner_email = f"master_owner_{TS}@test.com"
    r = client.post("/api/auth/register", json={"name": "Master Owner", "email": owner_email, "password": "Passw0rd1"})
    assert r.status_code == 200, r.text
    STATE["owner_token"] = r.json()["session_token"]
    STATE["owner_id"] = r.json()["user"]["user_id"]

    r = client.post("/api/auth/login", json={"email": owner_email, "password": "WrongPass1"})
    assert r.status_code == 401 and r.json()["detail"] == "Невалиден имейл или парола"

    r = client.post("/api/auth/login", json={"email": owner_email, "password": "Passw0rd1"})
    assert r.status_code == 200

    r = client.get("/api/auth/me", headers=h(STATE["owner_token"]))
    assert r.status_code == 200 and r.json()["role"] == "owner"


def test_invoices(client):
    owner = h(STATE["owner_token"])
    r = client.post("/api/invoices", headers=owner, json={
        "supplier": "Тест Доставчик ЕООД", "invoice_number": "INV-0001",
        "amount_without_vat": 100, "vat_amount": 20, "total_amount": 120, "date": "2026-09-01",
    })
    assert r.status_code == 200, r.text
    invoice_id = r.json()["id"]

    r = client.get("/api/invoices", headers=owner)
    assert r.status_code == 200 and len(r.json()) == 1

    r = client.get("/api/invoices?supplier=Тест", headers=owner)
    assert r.status_code == 200 and len(r.json()) == 1

    r = client.get(f"/api/invoices/{invoice_id}", headers=owner)
    assert r.status_code == 200

    r = client.put(f"/api/invoices/{invoice_id}", headers=owner, json={"total_amount": 150})
    assert r.status_code == 200 and r.json()["total_amount"] == 150

    # Re-saving the exact same value must not 404 (modified_count == 0 is
    # not "not found" - see the fix for this bug).
    r = client.put(f"/api/invoices/{invoice_id}", headers=owner, json={"total_amount": 150})
    assert r.status_code == 200

    # Regex-injection defense sanity check (should not error, should not
    # match unrelated data via an unescaped metacharacter).
    r = client.get("/api/invoices?supplier=.*", headers=owner)
    assert r.status_code == 200

    # Negative/zero amounts must be rejected at the API boundary.
    r = client.post("/api/invoices", headers=owner, json={
        "supplier": "X", "invoice_number": "BAD-1",
        "amount_without_vat": 10, "vat_amount": 2, "total_amount": -12, "date": "2026-09-01",
    })
    assert r.status_code == 422

    r = client.delete(f"/api/invoices/{invoice_id}", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/invoices", headers=owner)
    assert r.status_code == 200 and len(r.json()) == 0


def test_daily_revenue(client):
    owner = h(STATE["owner_token"])
    r = client.post("/api/daily-revenue", headers=owner, json={
        "date": "2026-09-10", "fiscal_revenue": 100, "pocket_money": 20, "vat_rate_percent": 20,
    })
    assert r.status_code == 200

    # Second save for the same date edits in place - REPLACED, not summed.
    r = client.post("/api/daily-revenue", headers=owner, json={
        "date": "2026-09-10", "fiscal_revenue": 100, "pocket_money": 15, "vat_rate_percent": 20,
    })
    assert r.status_code == 200
    assert r.json()["pocket_money"] == 15

    r = client.get("/api/daily-revenue/by-date/2026-09-10", headers=owner)
    assert r.status_code == 200 and r.json()["pocket_money"] == 15

    r = client.get("/api/daily-revenue/today", headers=owner)
    assert r.status_code == 200


def test_expenses(client):
    owner = h(STATE["owner_token"])
    r = client.post("/api/expenses", headers=owner, json={"description": "Гориво", "amount": 50, "date": "2026-09-10"})
    assert r.status_code == 200
    exp1_id = r.json()["id"]

    r = client.post("/api/expenses", headers=owner, json={"description": "Канцеларски материали", "amount": 30, "date": "2026-09-10"})
    assert r.status_code == 200

    r = client.get("/api/expenses", headers=owner)
    assert r.status_code == 200 and len(r.json()) == 2

    r = client.delete(f"/api/expenses/{exp1_id}", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/expenses", headers=owner)
    assert r.status_code == 200 and len(r.json()) == 1


def test_personal_expenses(client):
    owner = h(STATE["owner_token"])
    r = client.post("/api/personal-expenses", headers=owner, json={
        "amount": 200, "description": "Лична инвестиция", "expense_type": "investment", "category": "other",
        "period_month": 9, "period_year": 2026,
    })
    assert r.status_code == 200
    pe_id = r.json()["id"]

    r = client.get("/api/personal-expenses", headers=owner)
    assert r.status_code == 200 and len(r.json()["personal_expenses"]) == 1

    r = client.delete(f"/api/personal-expenses/{pe_id}", headers=owner)
    assert r.status_code == 200


def test_budget(client):
    owner = h(STATE["owner_token"])
    r = client.post("/api/budget", headers=owner, json={"month": CURRENT_MONTH, "expense_limit": 500, "alert_threshold": 80})
    assert r.status_code == 200

    # Re-saving for the same month overwrites, it doesn't duplicate.
    r = client.post("/api/budget", headers=owner, json={"month": CURRENT_MONTH, "expense_limit": 1000, "alert_threshold": 90})
    assert r.status_code == 200

    r = client.get("/api/budget", headers=owner)
    budgets = r.json()["budgets"]
    assert len(budgets) == 1 and budgets[0]["expense_limit"] == 1000

    r = client.get("/api/budget/status", headers=owner)
    assert r.status_code == 200 and r.json()["has_budget"] is True


def test_statistics(client):
    owner = h(STATE["owner_token"])
    r = client.get("/api/statistics/summary", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/statistics/chart-data?period=month", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/statistics/suppliers", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/statistics/items", headers=owner)
    assert r.status_code == 200


def test_roi_and_forecast(client):
    owner = h(STATE["owner_token"])
    r = client.get("/api/roi/analysis", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/roi/trend", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/forecast/expenses", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/forecast/revenue", headers=owner)
    assert r.status_code == 200


def test_employees_and_payroll(client):
    owner = h(STATE["owner_token"])
    r = client.post("/api/employees", headers=owner, json={
        "name": "Иван Иванов", "position": "Продавач", "base_salary": 1000, "agreement_type": "gross",
    })
    assert r.status_code == 200
    emp_id = r.json()["id"]

    r = client.get("/api/payroll/rates", headers=owner)
    assert r.status_code == 200

    r = client.post("/api/payroll/preview", headers=owner, json={
        "employee_id": emp_id, "period_month": 9, "period_year": 2026, "gross_amount": 1000,
    })
    assert r.status_code == 200 and "net_amount" in r.json()

    r = client.post("/api/payroll", headers=owner, json={
        "employee_id": emp_id, "period_month": 9, "period_year": 2026, "gross_amount": 1000,
    })
    assert r.status_code == 200
    payroll_id = r.json().get("id")

    r = client.get("/api/payroll", headers=owner)
    assert r.status_code == 200 and len(r.json()) >= 1

    if payroll_id:
        r = client.delete(f"/api/payroll/{payroll_id}", headers=owner)
        assert r.status_code == 200

    r = client.delete(f"/api/employees/{emp_id}", headers=owner)
    assert r.status_code == 200


def test_fixed_assets(client):
    owner = h(STATE["owner_token"])
    r = client.post("/api/assets", headers=owner, json={
        "name": "Лаптоп Dell", "category": "cat_iv", "acquisition_date": "2026-01-15",
        "in_service_date": "2026-01-15", "acquisition_value": 2000,
    })
    assert r.status_code == 200
    asset_id = r.json()["id"]

    r = client.get("/api/assets", headers=owner)
    assert r.status_code == 200 and len(r.json()) == 1

    r = client.get("/api/assets/categories", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/assets/summary", headers=owner)
    assert r.status_code == 200

    r = client.delete(f"/api/assets/{asset_id}", headers=owner)
    assert r.status_code == 200


def test_items_price_tracking(client):
    owner = h(STATE["owner_token"])
    r = client.get("/api/items/price-alert-settings", headers=owner)
    assert r.status_code == 200

    r = client.put("/api/items/price-alert-settings", headers=owner, json={"threshold_percent": 15, "enabled": True})
    assert r.status_code == 200

    r = client.get("/api/items/price-alerts", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/items/merge-mappings", headers=owner)
    assert r.status_code == 200


def test_notifications_and_company(client):
    owner = h(STATE["owner_token"])
    r = client.get("/api/notifications/settings", headers=owner)
    assert r.status_code == 200

    r = client.put("/api/notifications/settings", headers=owner, json={"vat_threshold_enabled": True, "vat_threshold_amount": 100})
    assert r.status_code == 200

    r = client.get("/api/company", headers=owner)
    assert r.status_code == 200


def test_utils(client):
    owner = h(STATE["owner_token"])
    r = client.get("/api/utils/validate-eik?eik=131059528", headers=owner)
    assert r.status_code == 200


def test_roles_invitations_permissions(client):
    owner = h(STATE["owner_token"])
    staff_email = f"master_staff_{TS}@test.com"
    r = client.post("/api/auth/register", json={"name": "Master Staff", "email": staff_email, "password": "Passw0rd1"})
    assert r.status_code == 200
    staff_token = r.json()["session_token"]
    staff_id = r.json()["user"]["user_id"]
    staff = h(staff_token)

    r = client.post("/api/invitations", headers=owner, json={"email": staff_email, "role": "accountant"})
    assert r.status_code == 200
    code = r.json()["invitation"]["code"]

    r = client.get("/api/invitations", headers=owner)
    assert r.status_code == 200 and len(r.json()) == 1

    r = client.post("/api/invitations/accept", headers=staff, json={"code": code})
    assert r.status_code == 200

    r = client.get("/api/auth/users", headers=owner)
    users = r.json()
    accountant_entry = next((u for u in users if u["user_id"] == staff_id), None)
    assert accountant_entry is not None
    assert accountant_entry["role"] == "accountant"
    assert "manage_budget" in accountant_entry.get("permissions", [])

    # Permission gating: accountant lacks add_revenue -> forbidden.
    r = client.post("/api/daily-revenue", headers=staff, json={"date": "2026-09-11", "fiscal_revenue": 10, "pocket_money": 0, "vat_rate_percent": 20})
    assert r.status_code == 403

    # accountant lacks view_off_book_expenses -> GET /expenses forbidden.
    r = client.get("/api/expenses", headers=staff)
    assert r.status_code == 403

    # accountant lacks view_pocket_money -> redacted (0), not blocked.
    r = client.get("/api/daily-revenue/by-date/2026-09-10", headers=staff)
    assert r.status_code == 200 and r.json()["pocket_money"] == 0
    r = client.get("/api/daily-revenue/by-date/2026-09-10", headers=owner)
    assert r.status_code == 200 and r.json()["pocket_money"] == 15

    r = client.get("/api/daily-revenue", headers=staff)
    staff_revenues = r.json()
    assert r.status_code == 200 and all(rv["pocket_money"] == 0 for rv in staff_revenues)
    r = client.get("/api/daily-revenue", headers=owner)
    owner_revenues = r.json()
    assert r.status_code == 200 and any(rv["pocket_money"] == 15 for rv in owner_revenues)

    # "Advanced statistics" is its own permission, not implied by
    # view_off_book_expenses or anything else an accountant gets by
    # default - staff registered with no role at all must be refused.
    plain_staff_email = f"master_plainstaff_{TS}@test.com"
    r = client.post("/api/auth/register", json={"name": "Plain Staff", "email": plain_staff_email, "password": "Passw0rd1"})
    assert r.status_code == 200
    plain_staff = h(r.json()["session_token"])
    r = client.post("/api/invitations", headers=owner, json={"email": plain_staff_email, "role": "staff"})
    assert r.status_code == 200
    plain_code = r.json()["invitation"]["code"]
    r = client.post("/api/invitations/accept", headers=plain_staff, json={"code": plain_code})
    assert r.status_code == 200
    r = client.get("/api/statistics/suppliers", headers=plain_staff)
    assert r.status_code == 403
    r = client.get("/api/items/price-alerts", headers=plain_staff)
    assert r.status_code == 200  # deliberately NOT gated - used by the invoice-detail badge flow

    r = client.get("/api/forecast/expenses", headers=plain_staff)
    assert r.status_code == 403  # plain staff lacks view_statistics
    r = client.get("/api/forecast/revenue", headers=plain_staff)
    assert r.status_code == 403

    # Financial-visibility redaction on /roi/analysis, /roi/trend and
    # /forecast/expenses: view_personal_investments/view_statistics gate the
    # endpoint itself, but pocket_money/off_book_expenses/profit stay
    # independently gated by their own view_* permission, same convention
    # as get_summary - zeroed out of every aggregate, not just blanked.
    r = client.post("/api/personal-expenses", headers=owner, json={
        "amount": 50, "description": "За тест на видимост", "expense_type": "investment", "category": "other",
        "period_month": 9, "period_year": 2026,
    })
    assert r.status_code == 200
    visibility_pe_id = r.json()["id"]

    r = client.put(f"/api/auth/role/{staff_id}", headers=owner, json={
        "role": "accountant",
        "permissions": [
            "view_audit_log", "manage_budget", "export_data", "view_statistics", "manage_invoices",
            "view_profit", "team_collaboration", "view_personal_investments",
        ],
    })
    assert r.status_code == 200

    r = client.get("/api/roi/analysis?month=9&year=2026", headers=owner)
    assert r.status_code == 200
    owner_roi = r.json()
    assert owner_roi["total_revenue"] == 115  # fiscal_revenue(100) + pocket_money(15)
    assert owner_roi["total_business_expense"] == 30  # off-book expense, no invoice left
    assert owner_roi["total_profit"] == 85

    r = client.get("/api/roi/analysis?month=9&year=2026", headers=staff)
    assert r.status_code == 200
    staff_roi = r.json()
    # No view_pocket_money/view_off_book_expenses: both folded in as zero,
    # not just blanked, so revenue/expense/profit stay internally coherent
    # for exactly what this viewer may see - not the owner's real numbers.
    assert staff_roi["total_revenue"] == 100
    assert staff_roi["total_business_expense"] == 0
    assert staff_roi["total_profit"] == 100  # has view_profit, so still shown
    assert staff_roi["total_profit"] != owner_roi["total_profit"]

    r = client.get("/api/roi/trend?months=2", headers=staff)
    assert r.status_code == 200
    assert r.json()["trend"][0]["profit"] == 100

    # Now strip view_profit too - profit/ROI must disappear from every
    # field AND from the ai_insights prose (no leaking the same figures via
    # text once the structured fields are null).
    r = client.put(f"/api/auth/role/{staff_id}", headers=owner, json={
        "role": "accountant",
        "permissions": [
            "view_audit_log", "manage_budget", "export_data", "view_statistics", "manage_invoices",
            "team_collaboration", "view_personal_investments",
        ],
    })
    assert r.status_code == 200

    r = client.get("/api/roi/analysis?month=9&year=2026", headers=staff)
    assert r.status_code == 200
    no_profit_roi = r.json()
    assert no_profit_roi["total_profit"] is None
    assert no_profit_roi["roi_percent"] is None
    assert no_profit_roi["is_profitable"] is None
    assert no_profit_roi["investment_covered"] is None
    assert no_profit_roi["ai_insights"] == ["🔒 Нямаш право да виждаш печалбата за избрания период"]

    r = client.get("/api/roi/trend?months=2", headers=staff)
    assert r.status_code == 200
    no_profit_trend = r.json()["trend"]
    assert no_profit_trend[0]["profit"] is None
    assert no_profit_trend[0]["roi_percent"] is None

    # The raw VAT-ledger Excel export must redact pocket_money too (same
    # right as the other views above), not just the aggregated endpoints.
    from io import BytesIO
    from openpyxl import load_workbook
    r = client.get("/api/export/vat-ledger/excel?start_date=2026-09-01&end_date=2026-09-30", headers=owner)
    assert r.status_code == 200
    owner_ws = load_workbook(BytesIO(r.content))["Дневник продажби"]
    assert owner_ws.cell(row=5, column=9).value == 15  # pocket_money, visible to the owner
    assert owner_ws.cell(row=5, column=10).value == 115  # total = fiscal_revenue + pocket_money

    r = client.get("/api/export/vat-ledger/excel?start_date=2026-09-01&end_date=2026-09-30", headers=staff)
    assert r.status_code == 200
    staff_ws = load_workbook(BytesIO(r.content))["Дневник продажби"]
    assert staff_ws.cell(row=5, column=9).value == 0  # no view_pocket_money -> redacted, not just blanked
    assert staff_ws.cell(row=5, column=10).value == 100

    # Off-book expense (30) must not leak into the expense-forecast
    # aggregate for a viewer without view_off_book_expenses either.
    r = client.get("/api/forecast/expenses", headers=owner)
    assert r.status_code == 200
    owner_hist = {e["month"]: e["amount"] for e in r.json()["historical"]}
    r = client.get("/api/forecast/expenses", headers=staff)
    assert r.status_code == 200
    staff_hist = {e["month"]: e["amount"] for e in r.json()["historical"]}
    assert owner_hist.get("2026-09", 0) - staff_hist.get("2026-09", 0) == 30

    r = client.delete(f"/api/personal-expenses/{visibility_pe_id}", headers=owner)
    assert r.status_code == 200

    # Owner removes the accountant - permissions must reset, not stay elevated.
    r = client.delete(f"/api/auth/users/{staff_id}", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/auth/me", headers=staff)
    me = r.json()
    assert r.status_code == 200 and me.get("role") in ("owner", "staff")


def test_backup_restore(client):
    owner = h(STATE["owner_token"])
    r = client.get("/api/backup/status", headers=owner)
    assert r.status_code == 200

    r = client.post("/api/backup/create", headers=owner)
    assert r.status_code == 200 and "invoices" in r.json()
    backup_payload = r.json()

    r = client.get("/api/backup/list", headers=owner)
    assert r.status_code == 200

    r = client.post("/api/backup/restore", headers=owner, json=backup_payload)
    assert r.status_code == 200


def test_exports(client):
    owner = h(STATE["owner_token"])
    r = client.get("/api/export/invoices/excel", headers=owner)
    assert r.status_code == 200 and "spreadsheet" in r.headers.get("content-type", "")

    r = client.get("/api/export/invoices/pdf", headers=owner)
    assert r.status_code == 200 and "pdf" in r.headers.get("content-type", "")

    r = client.get("/api/export/statistics/pdf", headers=owner)
    assert r.status_code == 200

    r = client.get("/api/export/vat-ledger/excel", headers=owner)
    assert r.status_code == 200


def test_ocr_size_gate(client):
    """Regression check: no AI key locally, so this only confirms an
    oversized image is rejected cleanly instead of crashing."""
    owner = h(STATE["owner_token"])
    huge_b64 = "A" * 15_000_000
    r = client.post("/api/ocr/scan", headers=owner, json={"image_base64": huge_b64})
    assert r.status_code in (413, 503)


def test_feedback_box(client):
    owner = h(STATE["owner_token"])
    r = client.post("/api/feedback", headers=owner, json={"message": "Много добро приложение, но искам PDF за дневника.", "is_anonymous": False})
    assert r.status_code == 200

    r = client.post("/api/feedback", headers=owner, json={"message": "Анонимен коментар за тест.", "is_anonymous": True})
    assert r.status_code == 200

    r = client.post("/api/feedback", headers=owner, json={"message": "a", "is_anonymous": False})
    assert r.status_code == 422

    r = client.get("/api/feedback", headers=owner)
    assert r.status_code == 403  # non-admin owner blocked from the feedback inbox


def test_scan_credits(client):
    """Fakturaplus's whole monetization surface - see ScanCreditService's
    module docstring. Uses its own freshly-registered owner/company so the
    balance assertions below aren't order-dependent on what other tests in
    this module may or may not have scanned."""
    import asyncio
    import server

    owner_email = f"master_credits_owner_{TS}@test.com"
    r = client.post("/api/auth/register", json={"name": "Credits Owner", "email": owner_email, "password": "Passw0rd1"})
    assert r.status_code == 200
    owner_token = r.json()["session_token"]
    owner_id = r.json()["user"]["user_id"]
    owner = h(owner_token)

    # Fresh company starts with the full free monthly quota and nothing purchased.
    r = client.get("/api/scan-credits", headers=owner)
    assert r.status_code == 200, r.text
    balance = r.json()
    assert balance["free_remaining"] == 10
    assert balance["free_quota"] == 10
    assert balance["purchased_balance"] == 0
    assert balance["total_remaining"] == 10
    assert balance["auto_reload_enabled"] is False

    # Fixed catalogue, matches the pricing conversation this was built from.
    r = client.get("/api/scan-credits/packages", headers=owner)
    assert r.status_code == 200
    packages = {p["id"]: p for p in r.json()}
    assert packages["small"]["total_scans"] == 10 and packages["small"]["price_eur"] == 2.50
    assert packages["medium"]["total_scans"] == 50 and packages["medium"]["price_eur"] == 11.00
    assert packages["large"]["total_scans"] == 150 and packages["large"]["price_eur"] == 30.00
    assert packages["business"]["total_scans"] == 400 and packages["business"]["price_eur"] == 72.00

    # No Stripe keys configured in the test environment -> a clear 503
    # instead of a crash, same disabled-until-configured pattern as the AI
    # features. Never a silent "pretend it worked".
    r = client.post("/api/scan-credits/purchase/package", headers=owner, json={"package_id": "small"})
    assert r.status_code == 503

    r = client.post("/api/scan-credits/auto-reload", headers=owner, json={"enabled": True, "package_id": "small"})
    assert r.status_code == 503

    r = client.post("/api/webhooks/stripe", json={})
    assert r.status_code == 503

    # manage_billing is owner-only, never configurable onto staff - see
    # ROLE_PERMISSIONS/ROLE_CONFIGURABLE_PERMISSIONS.
    staff_email = f"master_credits_staff_{TS}@test.com"
    r = client.post("/api/auth/register", json={"name": "Credits Staff", "email": staff_email, "password": "Passw0rd1"})
    assert r.status_code == 200
    staff = h(r.json()["session_token"])
    r = client.post("/api/invitations", headers=owner, json={"email": staff_email, "role": "staff"})
    assert r.status_code == 200
    code = r.json()["invitation"]["code"]
    r = client.post("/api/invitations/accept", headers=staff, json={"code": code})
    assert r.status_code == 200

    r = client.post("/api/scan-credits/purchase/package", headers=staff, json={"package_id": "small"})
    assert r.status_code == 403
    r = client.post("/api/scan-credits/auto-reload", headers=staff, json={"enabled": False})
    assert r.status_code == 403
    # Balance is shared company-wide though - any teammate can check it.
    r = client.get("/api/scan-credits", headers=staff)
    assert r.status_code == 200

    # Custom-quantity validation at the API boundary.
    r = client.post("/api/scan-credits/purchase/custom", headers=owner, json={"quantity": 0})
    assert r.status_code == 422
    r = client.post("/api/scan-credits/purchase/custom", headers=owner, json={"quantity": 1001})
    assert r.status_code == 422

    # --- Direct service-level tests: the part that actually protects money ---
    svc = server.scan_credit_service
    company_id = balance_company_id = None
    # Pull this fresh user's company_id the same way the endpoints do.
    me = client.get("/api/auth/me", headers=owner).json()
    company_id = me["company_id"]

    async def _run():
        # Stepped custom pricing matches the package brackets exactly.
        assert server.scan_credit_service_module.custom_price_eur(10) == 2.50
        assert server.scan_credit_service_module.custom_price_eur(11) == round(11 * (11.00 / 50), 2)
        assert server.scan_credit_service_module.custom_price_eur(50) == round(50 * (11.00 / 50), 2)
        assert server.scan_credit_service_module.custom_price_eur(51) == round(51 * (30.00 / 150), 2)
        assert server.scan_credit_service_module.custom_price_eur(1000) == round(1000 * (72.00 / 400), 2)

        # Spend down the free quota one at a time; the 11th spend with
        # nothing purchased must fail closed, never go negative.
        for _ in range(10):
            spent = await svc.spend_scan(company_id, owner_id, "Credits Owner")
            assert spent is True
        assert await svc.has_scans_remaining(company_id, owner_id) is False
        spent = await svc.spend_scan(company_id, owner_id, "Credits Owner")
        assert spent is False
        summary = await svc.get_balance_summary(company_id, owner_id)
        assert summary["free_remaining"] == 0 and summary["purchased_balance"] == 0

        # Credits only ever arrive via credit_purchase (the webhook path) -
        # never from spend_scan or any client-facing endpoint directly.
        await svc.credit_purchase(company_id, owner_id, "Credits Owner", scans=50, description="Пакет „Среден“", price_eur=11.0, payment_reference="pi_test_1")
        summary = await svc.get_balance_summary(company_id, owner_id)
        assert summary["purchased_balance"] == 50 and summary["total_remaining"] == 50

        # Free is exhausted, so the next spend must draw from purchased.
        spent = await svc.spend_scan(company_id, owner_id, "Credits Owner")
        assert spent is True
        summary = await svc.get_balance_summary(company_id, owner_id)
        assert summary["purchased_balance"] == 49

        # Webhook idempotency: the same Stripe event id is only ever
        # processed once, no matter how many times it's redelivered.
        assert await svc.mark_webhook_event_processed("evt_test_1") is True
        assert await svc.mark_webhook_event_processed("evt_test_1") is False
        assert await svc.mark_webhook_event_processed("evt_test_2") is True

        # History reflects both the spends and the purchase, newest first.
        history = await svc.get_history(company_id, owner_id)
        assert history[0]["type"] == "scan_used"
        assert any(tx["type"] == "purchase" and tx["delta"] == 50 for tx in history)

    asyncio.run(_run())

    r = client.get("/api/scan-credits/history", headers=owner)
    assert r.status_code == 200 and len(r.json()) >= 2

    r = client.get("/api/scan-credits/history/export", headers=owner)
    assert r.status_code == 200 and "spreadsheet" in r.headers.get("content-type", "")
